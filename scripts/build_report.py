"""Build the report Markdown and Word files from the saved experiment.

Run with Python 3.12 and the packages in requirements-report.txt. The builder
does not run retrieval or call Gemini. Export the Word file to PDF in a document
editor, or use a separately installed DOCX renderer. Rebuilding overwrites the
generated report files, so preserve personal revisions before running it again.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parents[1]
REPOSITORY = "https://github.com/varunpatil5050/rag-baseline-vs-reranking"


def read_json(relative: str):
    return json.loads((ROOT / relative).read_text(encoding="utf-8"))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-root", type=Path, default=ROOT)
    args = parser.parse_args()
    out = args.output_root.resolve()
    queries = {q["id"]: q for q in read_json("data/queries.json")}
    results = {r["case"]["id"]: r for r in read_json("results/retrieval_results.json")}
    answers = {(a["query_id"], a["variant"]): a for a in read_json("results/answers.json")}
    runtime = read_json("results/run_metadata.json")
    with (ROOT / "results/latency_summary.csv").open(newline="", encoding="utf-8") as handle:
        timings = {row["query_id"]: row for row in csv.DictReader(handle)}
    assert runtime["chunk_count"] == 409 and runtime["top_k"] == 10
    assert runtime["context_k"] == 3 and runtime["query_count"] == 5
    assert runtime["answers_reused_from_exact_request_cache"] == 10
    assert runtime["answers_generated_this_run"] == 0
    assert all(results[c]["designated_clause_recovered"] for c in ("D10", "D28", "D34"))

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Inches(0.64)
    section.left_margin = section.right_margin = Inches(0.68)
    section.header_distance = section.footer_distance = Inches(0.28)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
        style.element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
        for tag in ("w:pBdr", "w:shd"):
            for item in style.element.findall(".//" + qn(tag)):
                item.getparent().remove(item)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    for name, size, before, after in (
        ("Title", 18, 0, 7), ("Heading 1", 13, 10, 5), ("Heading 2", 11.5, 8, 4)
    ):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    doc.styles["List Bullet"].font.size = Pt(11)
    doc.styles["List Bullet"].paragraph_format.space_after = Pt(3)
    doc.styles["List Bullet"].paragraph_format.line_spacing = 1.05
    doc.core_properties.title = "RAG baseline retrieval and reranking"
    doc.core_properties.subject = "Comparison of dense retrieval and Cross-Encoder reranking"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = "RAG, Bi-Encoder, Cross-Encoder, NIST SP 800-63B-4"

    footer = section.footer.paragraphs[0]
    footer.style.font.name = "Arial"
    footer.style.font.size = Pt(9)
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Page ")
    run.font.name, run.font.size = "Arial", Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    field_run = OxmlElement("w:r")
    field_properties = OxmlElement("w:rPr")
    field_size = OxmlElement("w:sz")
    field_size.set(qn("w:val"), "18")
    field_properties.append(field_size)
    field_run.append(field_properties)
    field_text = OxmlElement("w:t")
    field_text.text = "1"
    field_run.append(field_text)
    field.append(field_run)
    footer._p.append(field)
    md: list[str] = []

    def para(text: str, *, bold_prefix: str | None = None, after: float | None = None):
        p = doc.add_paragraph()
        if bold_prefix:
            p.add_run(bold_prefix).bold = True
            p.add_run(text)
            prefix_space = bold_prefix[len(bold_prefix.rstrip()):]
            md.append(f"**{bold_prefix.rstrip()}**{prefix_space}{text}\n")
        else:
            p.add_run(text)
            md.append(text + "\n")
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    def heading(text: str, level: int = 1):
        doc.add_heading(text, level=level)
        md.append("#" * (level + 1) + " " + text + "\n")

    def link_paragraph(prefix: str, label: str, url: str, *, after: float = 5):
        p = doc.add_paragraph(prefix)
        p.paragraph_format.space_after = Pt(after)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
        r = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "000000")
        props.append(color)
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        props.append(underline)
        r.append(props)
        text = OxmlElement("w:t")
        text.text = label
        r.append(text)
        hyperlink.append(r)
        p._p.append(hyperlink)
        md.append(f"{prefix}[{label}]({url})\n")

    def answer(case: str, variant: str):
        label = "Answer 1 baseline" if variant == "baseline" else "Answer 2 reranked"
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        raw = answers[(case, variant)]["answer"]
        markdown_answer = re.sub(r"(?m)(^[^\n*].*)\n(?=\* )", r"\1\n\n", raw)
        md.append(f"**{label}**\n\n{markdown_answer}\n")
        # Preserve every response word and punctuation mark. Markdown bullet
        # notation is rendered as normal document bullets, as in a notebook.
        for block in raw.split("\n\n"):
            for line in block.splitlines():
                bullet = line.startswith("* ")
                text = line[2:] if bullet else line
                p = doc.add_paragraph(text, style="List Bullet" if bullet else None)
                p.paragraph_format.left_indent = Inches(0.12 if not bullet else 0.27)
                if bullet:
                    p.paragraph_format.first_line_indent = Inches(-0.15)
                p.paragraph_format.space_after = Pt(4)

    def case_heading(case: str, title: str, printed_page: int):
        heading(f"{case} {title}", level=2)
        q, r = queries[case], results[case]
        para(q["query"], bold_prefix="Question: ")
        para(
            f"PDF page {q['source_pdf_page']} (printed {printed_page}), section {q['source_section']}. "
            f"Chunk {r['tracked_evidence_chunk_id']}: rank {r['evidence_baseline_rank']} to "
            f"{r['evidence_reranked_rank']}.",
            after=5,
        )

    doc.add_paragraph("RAG baseline retrieval and reranking", style="Title")
    md.append("# RAG baseline retrieval and reranking\n")
    para(
        "Cross-Encoder reranking restored decisive evidence for three selected questions that baseline "
        "retrieval could not answer from its first three chunks. The experiment uses NIST SP 800-63B-4, "
        "Digital Identity Guidelines: Authentication and Authenticator Management, final July 2025 [1]. "
        "Two controls show why an improved rank need not improve an answer, and why missing candidates "
        "limit reranking."
    )
    link_paragraph("Repository: ", "github.com/varunpatil5050/rag-baseline-vs-reranking", REPOSITORY)

    heading("1 Architecture comparison")
    para(
        "The 129-page PDF yields 409 chunks from PDF pages 12-129, after excluding front matter. "
        "Each page is divided into 180-token windows with 40-token overlap. Both pipelines use this "
        "same corpus. The Bi-Encoder, sentence-transformers/all-MiniLM-L6-v2, embeds chunks and queries "
        "independently [2]. Chunk embeddings are precomputed. Both query and chunk vectors are "
        "normalized, so exact FAISS inner-product search is equivalent to cosine similarity."
    )
    para(
        "Baseline retrieval returns ten candidates and supplies its first three to Gemini. The "
        "Cross-Encoder, cross-encoder/ms-marco-MiniLM-L6-v2, jointly processes every one of the same "
        "ten query-chunk pairs, allowing self-attention across their tokens [3, 4]. It sorts candidates "
        "by descending relevance score and supplies the new first three. These relevance scores are "
        "not probabilities or cosine similarities."
    )
    para(
        "Both answers use gemini-3.8-flash, temperature 0, seed 42, and the same prompt template [5]. "
        "Only the three selected excerpts change. The prompt requires evidence-based answers, chunk "
        "citations, and acknowledgement of missing information; generation settings do not guarantee "
        "deterministic responses."
    )

    heading("2 Qualitative and rank shift analysis")
    para(
        "The five cases were selected from 44 exploratory queries. The three recoveries below were "
        "checked against the full baseline context for equivalent evidence. They are selected stress "
        "tests, not an estimate of average accuracy. The paired answers are reproduced in full."
    )
    case_heading("D10", "Resetting failed attempts", 22)
    answer("D10", "baseline")
    answer("D10", "reranked")
    para(
        "Baseline retrieved recovery-code material instead of the out-of-band counter restriction. "
        "The reranked context restores that restriction. The improvement is from an explicit "
        "abstention to a supported answer, not correction of a fabricated reset rule."
    )

    doc.add_page_break()
    case_heading("D28", "Reusing biometric samples", 34)
    answer("D28", "baseline")
    answer("D28", "reranked")
    para(
        "Baseline lacked the permission and consent clause. Reranking retrieves the local-use rule "
        "and the explicit-consent condition for other research. The cited provision does not state "
        "that template adaptation requires consent; this does not establish an exemption from all "
        "other applicable consent requirements."
    )
    case_heading("D34", "Sending a postal recovery notice", 52)
    answer("D34", "baseline")
    answer("D34", "reranked")
    para(
        "The baseline evidence covers recovery-code delivery, not the notification exception. "
        "Reranking restores the separate obligation to send a postal notice when the only other "
        "notification address received the issued recovery code."
    )

    doc.add_page_break()
    heading("Controls and retrieval limits", level=2)
    para(
        "D03 asks whether cloud-synced passkeys can satisfy AAL3. The designated prohibition, "
        "p020_c01 (PDF 20, printed 9, section 2.3.2), moves from rank 4 to 2. However, baseline rank 2 "
        "already contains equivalent evidence in p103_c02 (PDF 103, printed 92). Both answers correctly "
        "reject synced passkeys at AAL3, so this is not a failure recovery."
    )
    para(
        "D05 asks whether voice biometrics may be used. The prohibition in p043_c04 (PDF 43, printed "
        "32, section 3.2.3.2) is absent from the original ten candidates. Baseline incorrectly permits "
        "voice biometrics under general biometric conditions; the reranked answer abstains. Abstention "
        "avoids the unsupported permission, but the missing rule is not recovered. Full control "
        "answers and all candidate lists are retained in the notebook and results."
    )

    heading("3 Latency profiling")
    para(
        "Each query has one warm-up and seven measured repetitions, giving 35 samples. The run uses "
        "Python 3.12.3 on an ARM64 Mac with macOS 26.6.2, CPU inference, and one thread each for "
        "Torch and FAISS. "
        "The table reports medians in milliseconds from results/latency_summary.csv."
    )
    labels = ["Case", "Baseline\nms", "Reranking\nincrement ms", "Advanced\ntotal ms", "Total /\nbaseline"]
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.58, 1.14, 1.67, 1.48, 1.60]
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for i, label in enumerate(labels):
        table.rows[0].cells[i].text = label
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    md.append("| Case | Baseline median ms | Reranking increment median ms | Advanced total median ms | Total / baseline |")
    md.append("|---|---:|---:|---:|---:|")
    for case in ("D10", "D28", "D34", "D03", "D05"):
        row = timings[case]
        values = [case] + [f"{float(row[key]):.2f}" for key in (
            "baseline_median_ms", "rerank_increment_median_ms", "advanced_total_median_ms"
        )] + [f"{float(row['advanced_over_baseline']):.2f}x"]
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
        md.append("| " + " | ".join(values) + " |")
    md.append("")
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement(f"w:{edge}")
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "5")
                border.set(qn("w:color"), "D9D9D9")
                borders.append(border)
            tcpr.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge, value in (("top", 75), ("bottom", 75), ("left", 85), ("right", 85)):
                element = OxmlElement(f"w:{edge}")
                element.set(qn("w:w"), str(value))
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            tcpr.append(margins)
            if ri == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E9EEF2")
                tcpr.append(shade)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.size = Pt(10.5)
                    r.font.bold = ri == 0
    p = para(
        "Baseline includes query embedding, vector search, and candidate construction. Incremental "
        "reranking includes pair validation, tokenization, inference, sorting, and result construction. "
        "Advanced total adds both stages within each repetition; separately computed medians need "
        "not add exactly. Model loading, PDF processing, indexing, and Gemini generation are excluded."
    )
    p.paragraph_format.space_before = Pt(6)
    para(
        "Advanced retrieval is about 27.5-28.6 times slower in these measurements. Seven samples "
        "provide limited precision. The latest run recomputed retrieval and timings but reused all "
        "ten earlier Gemini responses from an exact prompt/model/configuration cache; it generated "
        "no new API answers. Original response timestamps remain in results/answers.json."
    )

    heading("4 Why reranking follows retrieval")
    para(
        "A Cross-Encoder must jointly process each query with each candidate chunk. Applying it "
        "first would require 409 pairs here rather than ten: 40.9 times as many pairs, not a measured "
        "40.9-fold latency increase. Bi-Encoder document vectors can be reused across queries, and "
        "the vector store narrows expensive joint scoring to ten candidates. This implementation "
        "uses exact FAISS search, not approximate search. Larger collections make the difference in "
        "pair counts grow, while D05 shows the cost of an initial retrieval miss. Reranking improves "
        "selection only within the candidates it receives."
    )
    heading("References", level=2)
    for label, url in (
        ("[1] NIST SP 800-63B-4 final publication", "https://csrc.nist.gov/pubs/sp/800/63/b/4/final"),
        ("[2] all-MiniLM-L6-v2 model card", "https://huggingface.co/sentence-transformers/all-MiniLM-L6-v2"),
        ("[3] MS MARCO MiniLM Cross-Encoder model card", "https://huggingface.co/cross-encoder/ms-marco-MiniLM-L6-v2"),
        ("[4] Sentence Transformers Cross-Encoder documentation", "https://www.sbert.net/docs/cross_encoder/usage/usage.html"),
        ("[5] Gemini content generation documentation", "https://ai.google.dev/gemini-api/docs/generate-content"),
    ):
        p = link_paragraph("", label, url, after=2)
    docx_path = out / "output/docx/RAG_Comparison_Report.docx"
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    markdown = "\n".join(line.rstrip() for line in "\n".join(md).splitlines())
    (out / "REPORT.md").write_text(markdown.rstrip() + "\n", encoding="utf-8")
    print(f"Created {out / 'REPORT.md'}")
    print(f"Created {docx_path}")
    print("Six recovery answers copied directly from results/answers.json.")


if __name__ == "__main__":
    main()
